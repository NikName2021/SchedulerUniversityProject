import pandas as pd
import json
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

WEEK = {0: "Понедельник", 1: "Вторник", 2: "Среда", 3: "Четверг", 4: "Пятница", 5: "Суббота", 6: "Воскресенье"}

COLORS = {
    "lec": "4472C4",  # Синий — лекция
    "sem": "70AD47",  # Зелёный — семинар
    "yellow": "FFD966", # Желтый — мягкая ошибка
    "red": "FF0000"     # Красный — жесткая ошибка
}

def save(final_schedule, warnings, SLOTS):
    if not final_schedule:
        print("❌ Нет решения для генерации расписания.")
        return

    rows = []
    for ev in final_schedule:
        date_obj = ev["slot"][0]
        wd = date_obj.weekday()
        subject_str = f'{ev["subject"]} ({ev["type"]})\nАуд: {ev["room"]}'
        
        rows.append({
            "date": str(date_obj),
            "weekday": WEEK[wd],
            "lesson": ev["slot"][1],
            "group": ev["group"],
            "subject": subject_str,
            "type": ev["type"],
            "warning": ev.get("warning", "")
        })

    df = pd.DataFrame(rows)
    df = df.sort_values(["date", "lesson", "group"])

    print("\n=== КОНФЛИКТЫ И ОШИБКИ ===")
    if not warnings:
        print("✅ Все идеально, конфликтов нет!")
    else:
        for w in warnings:
            if "date" in w and "lesson" in w:
                print(f"[{w['date']} Пара {w['lesson']}] {w['group']} - {w['subject']}: {w['msg']}")
            else:
                print(f"❌ НЕ ВЛЕЗЛО В РАСПИСАНИЕ: {w['group']} - {w['subject']} ({w.get('type', '')}) - {w.get('msg', 'Нет возможности поставить')}")

    # Сохраняем JSON отчет
    with open("result.json", "w", encoding="utf-8") as f:
        json.dump({"schedule": final_schedule, "warnings": warnings}, f, ensure_ascii=False, indent=4, default=str)

    # =============================
    # PIVOT ДЛЯ ТАБЛИЦЫ
    # =============================
    all_slots_idx = pd.MultiIndex.from_tuples(
        [(str(s[0]), WEEK[s[0].weekday()], s[1]) for s in SLOTS],
        names=["date", "weekday", "lesson"]
    )

    pivot_subject = df.pivot_table(index=["date", "weekday", "lesson"], columns="group", values="subject", aggfunc="first").reindex(all_slots_idx).reset_index()
    pivot_type = df.pivot_table(index=["date", "weekday", "lesson"], columns="group", values="type", aggfunc="first").reindex(all_slots_idx).reset_index()
    pivot_warning = df.pivot_table(index=["date", "weekday", "lesson"], columns="group", values="warning", aggfunc="first").reindex(all_slots_idx).reset_index()

    # =============================
    # ЭКСПОРТ В EXCEL С РАСКРАСКОЙ
    # =============================
    with pd.ExcelWriter("schedule.xlsx", engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="raw", index=False)
        ws = writer.book.create_sheet("pivot")

        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )

        header_row = 1
        data_start_row = 2
        data_start_col = 4 

        group_cols = pivot_type.columns[3:]
        current_row_idx = data_start_row
        prev_date = None

        for i in range(len(pivot_type)):
            row_type = pivot_type.iloc[i]
            row_subj = pivot_subject.iloc[i]
            row_warn = pivot_warning.iloc[i]
            
            current_date = row_type['date']
            if prev_date is not None and current_date != prev_date:
                for j in range(3 + len(group_cols)):
                    cell = ws.cell(row=current_row_idx, column=j + 1)
                    cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
                current_row_idx += 1

            prev_date = current_date

            for j in range(3):
                cell = ws.cell(row=current_row_idx, column=j + 1)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if j == 1: cell.font = Font(italic=True)
                
                if j == 0: cell.value = row_type['date']
                elif j == 1: cell.value = row_type['weekday']
                elif j == 2: cell.value = row_type['lesson']

            for j, group in enumerate(group_cols):
                c_type = row_type[group]
                c_warn = row_warn[group]
                c_subj = row_subj[group]
                
                cell = ws.cell(row=current_row_idx, column=data_start_col + j)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

                if pd.notna(c_type) and c_type in ["lec", "sem"]:
                    bg_color = COLORS[c_type]
                    text_color = "FFFFFF"
                    
                    if pd.notna(c_warn) and c_warn:
                        if "НЕТ АУДИТОРИИ" in str(c_warn):
                            bg_color = COLORS["red"]
                        else:
                            bg_color = COLORS["yellow"]
                            text_color = "000000"

                    cell.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")
                    cell.font = Font(color=text_color, bold=True)
                    
                    val = str(c_subj)
                    if pd.notna(c_warn) and c_warn:
                        val += f"\n⚠️ {c_warn}"
                    cell.value = val

            current_row_idx += 1

        for j, col_name in enumerate(pivot_type.columns):
            cell = ws.cell(row=header_row, column=j + 1)
            cell.value = {"date": "Дата", "weekday": "День", "lesson": "Пара"}.get(col_name, col_name)
            cell.fill = PatternFill(start_color="333333" if j >= 3 else "555555", end_color="333333" if j >= 3 else "555555", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 14
        ws.column_dimensions['C'].width = 8
        for j in range(len(group_cols)):
            ws.column_dimensions[chr(ord('D') + j)].width = 25

        legend_df = pd.DataFrame({
            "Цвет": ["🔵 Синий", "🟢 Зелёный", "🟡 Желтый", "🔴 Красный"],
            "Означает": ["Лекция (ОК)", "Семинар (ОК)", "Мягкая ошибка (Вместимость/Тип ауд.)", "Жесткая ошибка (Нет аудитории)"]
        })
        legend_df.to_excel(writer, sheet_name="Легенда", index=False)

    print("\nФайл сохранён: schedule.xlsx")

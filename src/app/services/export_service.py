import datetime
import io
import re

import pandas as pd
from core.constants import ALL_LESSONS
from database.all_models import GenerationTask, ScheduleEntry
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import select
from sqlalchemy.orm import selectinload

WEEK = {
    0: "Понедельник",
    1: "Вторник",
    2: "Среда",
    3: "Четверг",
    4: "Пятница",
    5: "Суббота",
    6: "Воскресенье",
}

COLORS = {
    "Лекция": "4472C4",
    "Семинар": "70AD47",
    "Лабораторная": "70AD47",
    "yellow": "FFD966",
    "red": "FF0000",
}

RAW_COLUMNS = ["date", "weekday", "lesson", "group", "subject", "type", "warning"]
UNASSIGNED_COLUMNS = [
    "task_id",
    "planning_week_id",
    "stream_id",
    "group",
    "subject",
    "type",
    "teacher",
    "missing_lessons",
    "target_lessons",
    "scheduled_lessons",
    "warning",
]
UNASSIGNED_WARNING = re.compile(r"Не выставлено\s+(\d+)\s+из\s+(\d+)\s+занятий")


def escape_excel_formula(value: str) -> str:
    """Keep imported text from being interpreted as an Excel formula."""

    if value and value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


async def generate_excel_report(
    session, task_id: int | None = None, semester_batch_id: str | None = None
):
    stmt = select(ScheduleEntry).options(selectinload(ScheduleEntry.teacher))
    if task_id is not None:
        stmt = stmt.filter(ScheduleEntry.task_id == task_id)
    elif semester_batch_id is not None:
        stmt = stmt.join(
            GenerationTask, GenerationTask.id == ScheduleEntry.task_id
        ).where(GenerationTask.semester_batch_id == semester_batch_id)
    else:
        stmt = stmt.join(
            GenerationTask, GenerationTask.id == ScheduleEntry.task_id
        ).where(GenerationTask.publication_status == "published")

    stmt = stmt.order_by(ScheduleEntry.date, ScheduleEntry.lesson_number)
    result = await session.execute(stmt)
    entries = result.scalars().all()

    if not entries:
        return None

    rows = []
    unassigned: dict[tuple, dict] = {}
    for ev in entries:
        if ev.date is None:
            warning = escape_excel_formula(ev.warning or "")
            key = (
                ev.task_id,
                ev.planning_week_id,
                ev.source_stream_id,
                ev.group_name,
                ev.event_name,
                ev.stream_type,
                ev.teacher.name if ev.teacher else "",
                warning,
            )
            if key not in unassigned:
                match = UNASSIGNED_WARNING.search(ev.warning or "")
                missing = int(match.group(1)) if match else 0
                target = int(match.group(2)) if match else None
                unassigned[key] = {
                    "task_id": ev.task_id,
                    "planning_week_id": ev.planning_week_id,
                    "stream_id": ev.source_stream_id,
                    "group": escape_excel_formula(ev.group_name),
                    "subject": escape_excel_formula(ev.event_name),
                    "type": escape_excel_formula(ev.stream_type),
                    "teacher": escape_excel_formula(
                        ev.teacher.name if ev.teacher else ""
                    ),
                    "missing_lessons": missing,
                    "target_lessons": target,
                    "scheduled_lessons": (
                        target - missing if target is not None else None
                    ),
                    "warning": warning,
                    "occurrences": 0,
                }
            unassigned[key]["occurrences"] += 1
            continue
        wd = ev.date.weekday()
        subject_str = escape_excel_formula(
            f"{ev.event_name} ({ev.stream_type})\n"
            f"{ev.teacher.name if ev.teacher else ''}\nАуд: {ev.room_id or ''}"
        )

        rows.append(
            {
                "date": ev.date.strftime("%Y-%m-%d"),
                "weekday": WEEK[wd],
                "lesson": ev.lesson_number,
                "group": escape_excel_formula(ev.group_name),
                "subject": subject_str,
                "type": escape_excel_formula(ev.stream_type),
                "warning": escape_excel_formula(ev.warning or ""),
            }
        )

    if not rows and not unassigned:
        return None

    df = pd.DataFrame(rows, columns=RAW_COLUMNS)
    unassigned_rows = []
    for item in unassigned.values():
        if not item["missing_lessons"]:
            item["missing_lessons"] = item["occurrences"]
        item.pop("occurrences")
        unassigned_rows.append(item)
    unassigned_rows.sort(
        key=lambda item: (
            item["task_id"] or 0,
            item["planning_week_id"] or 0,
            str(item["group"]),
            str(item["subject"]),
        )
    )
    unassigned_df = pd.DataFrame(unassigned_rows, columns=UNASSIGNED_COLUMNS)

    all_dates = sorted(df["date"].unique()) if rows else []

    slots_tuples = []
    for d_str in all_dates:
        d_obj = datetime.datetime.strptime(d_str, "%Y-%m-%d")
        for lesson_number in ALL_LESSONS:
            slots_tuples.append((d_str, WEEK[d_obj.weekday()], lesson_number))

    all_slots_idx = pd.MultiIndex.from_tuples(
        slots_tuples, names=["date", "weekday", "lesson"]
    )

    if rows:
        pivot_subject = (
            df.pivot_table(
                index=["date", "weekday", "lesson"],
                columns="group",
                values="subject",
                aggfunc="first",
            )
            .reindex(all_slots_idx)
            .reset_index()
        )
        pivot_type = (
            df.pivot_table(
                index=["date", "weekday", "lesson"],
                columns="group",
                values="type",
                aggfunc="first",
            )
            .reindex(all_slots_idx)
            .reset_index()
        )
        pivot_warning = (
            df.pivot_table(
                index=["date", "weekday", "lesson"],
                columns="group",
                values="warning",
                aggfunc="first",
            )
            .reindex(all_slots_idx)
            .reset_index()
        )
    else:
        empty_columns = ["date", "weekday", "lesson"]
        pivot_subject = pd.DataFrame(columns=empty_columns)
        pivot_type = pd.DataFrame(columns=empty_columns)
        pivot_warning = pd.DataFrame(columns=empty_columns)

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="raw", index=False)
        unassigned_df.to_excel(writer, sheet_name="Невыставленные", index=False)
        ws = writer.book.create_sheet("Расписание")

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        header_row = 1
        data_start_row = 2
        data_start_col = 4

        group_cols = pivot_type.columns[3:]
        current_row_idx = data_start_row
        prev_date = None

        for i in range(len(pivot_type)):
            row_type = pivot_type.iloc[i]
            row_subj = pivot_subject.iloc[i]
            row_warn = pivot_warning.iloc[i]

            current_date = row_type["date"]

            if prev_date is not None and current_date != prev_date:
                for j in range(3 + len(group_cols)):
                    cell = ws.cell(row=current_row_idx, column=j + 1)
                    cell.fill = PatternFill(
                        start_color="DDDDDD", end_color="DDDDDD", fill_type="solid"
                    )
                current_row_idx += 1

            prev_date = current_date

            for j in range(3):
                cell = ws.cell(row=current_row_idx, column=j + 1)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if j == 0:
                    cell.value = row_type["date"]
                elif j == 1:
                    cell.value = row_type["weekday"]
                elif j == 2:
                    cell.value = row_type["lesson"]

            for j, group in enumerate(group_cols):
                c_type = row_type[group]
                c_warn = row_warn[group]
                c_subj = row_subj[group]

                cell = ws.cell(row=current_row_idx, column=data_start_col + j)
                cell.border = thin_border
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )

                if pd.notna(c_type):
                    bg_color = COLORS.get(c_type, "CCCCCC")
                    text_color = "FFFFFF"

                    if pd.notna(c_warn) and c_warn:
                        if "Нет аудитории" in str(c_warn):
                            bg_color = COLORS["red"]
                        else:
                            bg_color = COLORS["yellow"]
                            text_color = "000000"

                    cell.fill = PatternFill(
                        start_color=bg_color, end_color=bg_color, fill_type="solid"
                    )
                    cell.font = Font(color=text_color, bold=True)

                    val = str(c_subj)
                    if pd.notna(c_warn) and c_warn:
                        val += f"\n⚠️ {c_warn}"
                    cell.value = escape_excel_formula(val)

            current_row_idx += 1

        for j, col_name in enumerate(pivot_type.columns):
            cell = ws.cell(row=header_row, column=j + 1)
            cell.value = escape_excel_formula(
                str(
                    {"date": "Дата", "weekday": "День", "lesson": "Пара"}.get(
                        col_name, col_name
                    )
                )
            )
            cell.fill = PatternFill(
                start_color="333333", end_color="333333", fill_type="solid"
            )
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 8

        unassigned_ws = writer.book["Невыставленные"]
        unassigned_ws.freeze_panes = "A2"
        unassigned_ws.auto_filter.ref = unassigned_ws.dimensions
        for cell in unassigned_ws[1]:
            cell.fill = PatternFill(
                start_color="C00000", end_color="C00000", fill_type="solid"
            )
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        unassigned_widths = {
            "A": 12,
            "B": 18,
            "C": 12,
            "D": 24,
            "E": 52,
            "F": 16,
            "G": 34,
            "H": 18,
            "I": 16,
            "J": 20,
            "K": 36,
        }
        for column, width in unassigned_widths.items():
            unassigned_ws.column_dimensions[column].width = width
        for row in unassigned_ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = thin_border

    output.seek(0)
    return output

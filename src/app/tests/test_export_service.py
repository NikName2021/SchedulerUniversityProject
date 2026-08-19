import datetime
import io

import pytest
from database import GenerationTask, ScheduleEntry, Teacher
from openpyxl import load_workbook
from services.export_service import generate_excel_report


@pytest.mark.asyncio
async def test_excel_export_contains_seventh_lesson_and_escapes_formulas(
    db_session,
) -> None:
    teacher = Teacher(name="@Преподаватель")
    task = GenerationTask(status="success")
    db_session.add_all([teacher, task])
    await db_session.flush()
    db_session.add(
        ScheduleEntry(
            task_id=task.id,
            group_name="=ГР-1",
            event_name="=HYPERLINK(\"https://example.invalid\")",
            stream_type="Лекция",
            teacher_id=teacher.id,
            room_id="+101",
            date=datetime.datetime(2026, 9, 7),
            lesson_number=7,
            warning="@warning",
        )
    )
    await db_session.commit()

    output = await generate_excel_report(db_session, task_id=task.id)

    assert output is not None
    workbook = load_workbook(io.BytesIO(output.getvalue()), data_only=False)
    raw = workbook["raw"]
    schedule = workbook["Расписание"]
    assert all(
        cell.data_type != "f"
        for worksheet in (raw, schedule)
        for row in worksheet.iter_rows()
        for cell in row
    )
    assert any(cell.value == 7 for cell in schedule["C"])
    assert any(
        isinstance(cell.value, str) and cell.value.startswith("'=ГР-1")
        for cell in schedule[1]
    )


@pytest.mark.asyncio
async def test_excel_export_with_only_unassigned_entries_adds_summary_sheet(
    db_session,
) -> None:
    task = GenerationTask(status="partial")
    db_session.add(task)
    await db_session.flush()
    db_session.add_all(
        [
            ScheduleEntry(
                task_id=task.id,
                group_name="ГР-1",
                event_name="Физика",
                stream_type="Семинар",
                warning="Не выставлено 2 из 4 занятий",
            ),
            ScheduleEntry(
                task_id=task.id,
                group_name="ГР-1",
                event_name="Физика",
                stream_type="Семинар",
                warning="Не выставлено 2 из 4 занятий",
            ),
        ]
    )
    await db_session.commit()

    output = await generate_excel_report(db_session, task_id=task.id)

    assert output is not None
    workbook = load_workbook(io.BytesIO(output.getvalue()), data_only=False)
    assert workbook.sheetnames == ["raw", "Невыставленные", "Расписание"]
    headers = [cell.value for cell in workbook["Невыставленные"][1]]
    values = [cell.value for cell in workbook["Невыставленные"][2]]
    unassigned = dict(zip(headers, values))
    assert unassigned == {
        "task_id": task.id,
        "planning_week_id": None,
        "stream_id": None,
        "group": "ГР-1",
        "subject": "Физика",
        "type": "Семинар",
        "teacher": None,
        "missing_lessons": 2,
        "target_lessons": 4,
        "scheduled_lessons": 2,
        "warning": "Не выставлено 2 из 4 занятий",
    }
    assert all(
        cell.data_type != "f"
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
    )
